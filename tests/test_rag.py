import io
import zipfile

import pytest
from docx import Document

from app.core.config import Settings
from app.rag import DemoHashEmbedding, KnowledgeBases, read_document
from app.rag.embedding_cache import CachedEmbedding


def test_txt_ingest_search_deduplicate_restart_delete(settings):
    bases = KnowledgeBases(settings)
    kb = bases.for_tenant("a")
    raw = "Thời hạn hoàn cước là 7 ngày. Hồ sơ gồm hóa đơn.".encode()
    record = kb.ingest("hoan_cuoc.txt", raw)
    hits = kb.retrieve("thời hạn hoàn cước")
    assert hits and "7 ngày" in hits[0]["excerpt"]
    assert hits[0]["filename"] == "hoan_cuoc.txt"
    assert kb.ingest("another.txt", raw)["duplicate"]
    assert len(kb.snapshot.records) == 1
    assert KnowledgeBases(settings).for_tenant("a").retrieve("hoàn cước")
    assert kb.remove(record["id"])
    assert not kb.retrieve("hoàn cước")
    assert not KnowledgeBases(settings).for_tenant("a").snapshot.nodes


def test_docx_preserves_tables_and_locations(settings):
    document = Document()
    document.add_heading("Chính sách mẫu", 1)
    document.add_paragraph("Mã bưu gửi được kiểm tra trước khi trả lời khách hàng.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Dịch vụ", "Thời hạn"
    table.cell(1, 0).text, table.cell(1, 1).text = "Hoàn cước", "7 ngày"
    buffer = io.BytesIO()
    document.save(buffer)
    docs = read_document("policy.docx", buffer.getvalue(), "docid", settings)
    assert any("bảng 1, hàng 2" in d.metadata["location"] and "7 ngày" in d.text for d in docs)
    assert all("trang" not in d.metadata["location"] for d in docs)
    kb = KnowledgeBases(settings).for_tenant("a")
    kb.ingest("policy.docx", buffer.getvalue())
    assert any("7 ngày" in h["excerpt"] for h in kb.retrieve("hoàn cước thời hạn"))


@pytest.mark.parametrize(
    "filename,raw",
    [
        ("x.doc", b"data"),
        ("x.dox", b"data"),
        ("x.exe", b"data"),
        ("x.txt", b"\x00"),
        ("x.txt", b"\xff"),
        ("x.txt", b"   "),
        ("x.docx", b"bad zip"),
    ],
)
def test_rejects_invalid_uploads(settings, filename, raw):
    with pytest.raises(ValueError):
        read_document(filename, raw, "x", settings)


def test_docx_expansion_limit(settings):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * 2000)
    small = settings.model_copy(update={"max_docx_expanded_bytes": 1024})
    with pytest.raises(ValueError, match="giải nén"):
        read_document("x.docx", buffer.getvalue(), "x", small)


def test_tenant_isolation_filename_traversal_and_no_match(settings):
    bases = KnowledgeBases(settings)
    a, b = bases.for_tenant("a"), bases.for_tenant("b")
    a.ingest("../../secret.txt", "Thời hạn hoàn cước là 7 ngày".encode())
    assert a.snapshot.records[0]["filename"] == "secret.txt"
    assert not b.retrieve("hoàn cước")
    assert not a.retrieve("quasar neutrino gamma")


def test_config_change_requires_reindex(settings):
    KnowledgeBases(settings).for_tenant("a").ingest("x.txt", "Hoàn cước 7 ngày".encode())
    changed = settings.model_copy(update={"chunk_size": 256})
    with pytest.raises(ValueError, match="reindex"):
        KnowledgeBases(changed).for_tenant("a")


def test_failed_rebuild_does_not_publish_partial_document(settings, monkeypatch):
    kb = KnowledgeBases(settings).for_tenant("a")
    kb.ingest("good.txt", "Hoàn cước 7 ngày".encode())
    before = kb.snapshot.revision
    monkeypatch.setattr(
        kb, "_commit", lambda *args: (_ for _ in ()).throw(RuntimeError("provider failed"))
    )
    with pytest.raises(RuntimeError):
        kb.ingest("new.txt", "Doanh thu 120 triệu đồng".encode())
    assert kb.snapshot.revision == before and len(kb.snapshot.records) == 1
    assert len(list(kb.raw_dir.iterdir())) == 1


def test_embedding_cache_reuses_vectors_and_separates_query_from_document(tmp_path):
    class CountingEmbedding(DemoHashEmbedding):
        calls: int = 0

        def _get_text_embedding(self, text):
            self.calls += 1
            return super()._get_text_embedding(text)

    inner = CountingEmbedding()
    cache = CachedEmbedding(inner, tmp_path / "cache.sqlite3", "m1")
    cache.get_text_embedding_batch(["hoàn cước", "hoàn cước"])
    assert inner.calls == 1
    cache.get_text_embedding_batch(["hoàn cước", "doanh thu"])
    assert inner.calls == 2
    reloaded = CachedEmbedding(inner, tmp_path / "cache.sqlite3", "m1")
    reloaded.get_text_embedding("hoàn cước")
    assert inner.calls == 2
    with __import__("sqlite3").connect(tmp_path / "cache.sqlite3") as db:
        before = db.execute("SELECT COUNT(*) FROM vectors").fetchone()[0]
    reloaded.get_query_embedding("hoàn cước")
    with __import__("sqlite3").connect(tmp_path / "cache.sqlite3") as db:
        assert db.execute("SELECT COUNT(*) FROM vectors").fetchone()[0] == before


def test_live_mode_rejects_demo_embedding_and_missing_keys(tmp_path):
    with pytest.raises(ValueError):
        Settings(_env_file=None, demo_mode=False, data_dir=tmp_path)
    with pytest.raises(ValueError, match="demo-only"):
        Settings(
            _env_file=None,
            demo_mode=False,
            data_dir=tmp_path,
            api_keys_json='{"1234567890123456":"a"}',
            llm_api_key="dummy",
            llm_model="real",
        )


def test_bkai_chunk_respects_encoder_window(tmp_path):
    with pytest.raises(ValueError, match="CHUNK_SIZE"):
        Settings(
            _env_file=None,
            data_dir=tmp_path,
            embedding_provider="huggingface",
            local_embedding_model="bkai-foundation-models/vietnamese-bi-encoder",
            chunk_size=257,
            embedding_max_length=256,
        )


def test_reranker_reorders_only_the_hybrid_candidate_pool(settings):
    class StubReranker:
        def rerank(self, _query, candidates):
            reordered = list(reversed(candidates))
            return [
                {**item, "reranker_score": float(index)} for index, item in enumerate(reordered, 1)
            ]

    rerank_settings = settings.model_copy(
        update={"reranker_enabled": True, "reranker_candidate_k": 6}
    )
    kb = KnowledgeBases(rerank_settings).for_tenant("a")
    kb.ingest("first.txt", "Hoàn cước cần hóa đơn hợp lệ để xử lý.".encode())
    kb.ingest("second.txt", "Hoàn cước được giải quyết trong 7 ngày làm việc.".encode())
    kb.reranker = StubReranker()

    hits = kb.retrieve("hoàn cước", top_k=2)

    assert len(hits) == 2
    assert all("reranker_score" in hit for hit in hits)


def test_metadata_filtering_global_and_session_isolation(settings):
    kb = KnowledgeBases(settings).for_tenant("test-tenant")
    # 1. Global system document
    kb.ingest(
        "global_policy.txt",
        "Thời hạn hoàn cước hệ thống VNPost là 7 ngày.".encode(),
        is_global=True,
    )
    # 2. Session A personal document
    kb.ingest(
        "user_a_contract.txt",
        "Hợp đồng riêng khách hàng A mã số HDA-100.".encode(),
        is_global=False,
        session_id="session-a",
    )
    # 3. Session B personal document
    kb.ingest(
        "user_b_contract.txt",
        "Hợp đồng riêng khách hàng B mã số HDB-200.".encode(),
        is_global=False,
        session_id="session-b",
    )

    # Session A query should retrieve global doc and session A doc, but NEVER session B doc
    hits_a_global = kb.retrieve("hoàn cước", session_id="session-a")
    assert any(h["filename"] == "global_policy.txt" for h in hits_a_global)

    hits_a_contract = kb.retrieve("Hợp đồng", session_id="session-a")
    assert any(h["filename"] == "user_a_contract.txt" for h in hits_a_contract)
    assert not any(h["filename"] == "user_b_contract.txt" for h in hits_a_contract)

    # Session B query should retrieve global doc and session B doc, but NEVER session A doc
    hits_b_contract = kb.retrieve("Hợp đồng", session_id="session-b")
    assert any(h["filename"] == "user_b_contract.txt" for h in hits_b_contract)
    assert not any(h["filename"] == "user_a_contract.txt" for h in hits_b_contract)

    # Anonymous / General query should not retrieve private session docs
    hits_general = kb.retrieve("Hợp đồng", session_id=None)
    assert not any(
        h["filename"] in ("user_a_contract.txt", "user_b_contract.txt") for h in hits_general
    )


def test_tokens_unique_no_accents():
    from app.rag import tokens

    # Từ "tính" chỉ được trả về 1 token là "tinh", không được trả về cả "tính" và "tinh"
    toks = tokens("tính")
    assert toks == ["tinh"]
    assert len(toks) == 1


def test_filter_relevant_hits_min_overlap():
    from app.rag import filter_relevant_hits

    # Với query 4 từ khóa: "quy trình hoàn hàng" -> q_terms = {"quy", "trinh", "hoan", "hang"}
    # Ngưỡng min_overlap = 2.
    # Excerpt chỉ chứa 1 từ trùng ("quy") không được lọt qua cổng.
    hits = [
        {"excerpt": "quy định chung", "filename": "1.txt"},
        {"excerpt": "quy trình hoàn hàng bưu chính", "filename": "2.txt"},
    ]
    res = filter_relevant_hits("quy trình hoàn hàng", hits)
    assert len(res) == 1
    assert res[0]["filename"] == "2.txt"


def test_resolve_embedding_device(monkeypatch):
    from app.rag.engine import resolve_embedding_device

    assert resolve_embedding_device("cpu") == "cpu"
    assert resolve_embedding_device("cuda") == "cuda"

    # When auto without cuda available
    import sys

    fake_torch = type(sys)("torch")
    fake_torch.cuda = type(sys)("cuda")
    fake_torch.cuda.is_available = lambda: False
    monkeypatch.setitem(sys.modules, "torch", fake_torch)
    assert resolve_embedding_device("auto") == "cpu"

    # When auto with cuda available
    fake_torch.cuda.is_available = lambda: True
    assert resolve_embedding_device("auto") == "cuda"


def test_bm25_term_frequency_preservation(settings):
    bases = KnowledgeBases(settings)
    kb = bases.for_tenant("tf_tenant")
    # File 1 has single mention of 'khiếu nại'
    kb.ingest(
        "doc1.txt",
        "Quy định về thời hạn tiếp nhận và thụ lý hồ sơ khiếu nại của khách hàng.".encode(),
    )
    # File 2 has dense repeated mentions of 'khiếu nại'
    kb.ingest(
        "doc2.txt",
        "Khiếu nại khiếu nại khiếu nại dịch vụ bưu chính được giải quyết qua cổng chăm sóc khách hàng.".encode(),
    )

    # Check BM25 score directly: doc2 with 3 mentions must score higher than doc1 with 1 mention
    scores = kb.snapshot.bm25.get_scores(["khieu", "nai"])
    assert scores[1] > scores[0]


def test_session_scoped_ingest_deduplication_isolation(settings):
    bases = KnowledgeBases(settings)
    kb = bases.for_tenant("scoped_tenant")
    raw = "Hợp đồng dịch vụ riêng tư số 999.".encode()

    # Session A uploads
    rec_a = kb.ingest("contract.txt", raw, session_id="sess_a")
    assert rec_a["duplicate"] is False
    assert rec_a["session_id"] == "sess_a"

    # Session B uploads identical bytes
    rec_b = kb.ingest("contract.txt", raw, session_id="sess_b")
    # Should be isolated to Session B, not treated as duplicate of Session A
    assert rec_b["duplicate"] is False
    assert rec_b["session_id"] == "sess_b"
    assert rec_a["id"] != rec_b["id"]

    # Session B can retrieve its own document
    hits_b = kb.retrieve("Hợp đồng dịch vụ riêng tư", session_id="sess_b")
    assert any(h["document_id"] == rec_b["id"] for h in hits_b)

    # Session A deletes its session
    deleted = kb.remove_session("sess_a")
    assert deleted == 1

    # Session B's document is still intact and retrievable
    hits_b_after = kb.retrieve("Hợp đồng dịch vụ riêng tư", session_id="sess_b")
    assert any(h["document_id"] == rec_b["id"] for h in hits_b_after)
